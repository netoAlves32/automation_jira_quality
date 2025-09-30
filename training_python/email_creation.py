import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


CZ = "CZ SUROESTE"
ENLACE = "HARVIN"

# Cargar datos
df = pd.read_excel("/Users/neto_shayo/Desktop/NETO/ICBF/SIMAT/Alertas_SIMAT_31082025.xlsx", sheet_name="Hoja1", engine="openpyxl")

# Filtrar por CZ y ENLACE_CUENTAME = HARVIN
df_cz_harvin = df[(df["enlaceCuentame"] == ENLACE) & (df["CentroZonalUDS"] == CZ)]

# Obtener todas las EAS únicas
eas_unicas = df_cz_harvin["EntidadContratista"].unique()

# Columnas de la tabla
columns = [
    "NumeroDocumentoBeneficiario", "PrimerNombreBeneficiario", "SegundoNombreBeneficiario",
    "PrimerApellidoBeneficiario", "SegundoApellidoBeneficiario", "FechaNacimientoBeneficiario",
    "EdadAños_FechaBackup"
]

# Generar un documento por cada EAS
for eas in eas_unicas:
    doc = Document()
    # Encabezado general
    header = doc.add_paragraph()
    header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = header.add_run(f"{eas}\nCentro Zonal: {CZ}")
    run.bold = True
    run.font.size = Pt(14)
    doc.add_paragraph("\nCordial saludo,\n")
    doc.add_paragraph(
        f"Desde la Regional Antioquia, en el marco de la implementación del sistema de alertas frente al registro de información para la presente vigencia, reenvío la siguiente información. "
        f"El pasado 31 de agosto se emitió una alerta relacionada con el SIMAT (Sistema Integrado de Matrícula), herramienta informática del Ministerio de Educación Nacional de Colombia, "
        f"cuya función principal es organizar y controlar el proceso de matrícula en las instituciones educativas oficiales, desde la inscripción hasta la asignación de cupos. Esta alerta fue generada por el Ministerio, con el objetivo de identificar a los niños, niñas y adolescentes (NNA) que presentan concurrencia en la atención por parte de diferentes entidades del Estado.\n"
        f"Solicito de manera urgente el envío de una carta con LOGO incluido de la EAS correspondiente {eas} "
        f"adscritas al Centro Zonal: {CZ}, en la que se indique de manera clara su deseo de continuar con la atención por parte del ICBF "
        "firmada por parte de los acudientes de los siguientes beneficiario/a(s):\n"
    )
    uds_unicas = df_cz_harvin[df_cz_harvin["EntidadContratista"] == eas]["UnidadServicio"].unique()
    for uds in uds_unicas:
        p_uds = doc.add_paragraph()
        run_uds = p_uds.add_run(f"\nUnidad de Servicio: {uds}")
        run_uds.bold = True
        filtered_df = df_cz_harvin[(df_cz_harvin["UnidadServicio"] == uds) & (df_cz_harvin["EntidadContratista"] == eas)]
        if filtered_df.empty:
            doc.add_paragraph("No hay beneficiarios para esta UDS.")
            continue
        table = doc.add_table(rows=1, cols=len(columns))
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(columns):
            hdr_cells[i].text = col
        for _, row in filtered_df.iterrows():
            row_cells = table.add_row().cells
            for i, col in enumerate(columns):
                value = row[col]
                row_cells[i].text = "" if pd.isna(value) else str(value)
        doc.add_paragraph("")
    # Párrafo de cierre
    doc.add_paragraph(
        "\nLa carta debe indicar:\n"
        "• Si el beneficiario continuará recibiendo atención por parte del ICBF, explicando causas.\n"
        "• Cuando envíen el archivo de las cartas debe ser por cada beneficiario y deben de nombrarlo con el número de identidad\n"
        "• Si el beneficiario no continuará, manifestando su decisión para proceder en la desvinculación.\n\n"
        "Sin este documento, no será posible garantizar la permanencia del menor en el programa y, en cumplimiento de lo acordado con el Ministerio de Educación y demás entidades, "
        "se procederá a su desvinculación para evitar pagos duplicados. Agradecemos su pronta gestión para evitar contratiempos.\n\n"
        "Agradezco su colaboración y envío de la información solicitada con la mayor celeridad posible.\n"
        "Quedo atento a su pronta respuesta.\n\n"
        "Cordialmente,\n"
    )
    import os
    base_dir = "/Users/neto_shayo/Desktop/NETO/ICBF/SIMAT"
    cz_dir = os.path.join(base_dir, CZ.replace(' ', '_'), "correos")
    os.makedirs(cz_dir, exist_ok=True)
    output_path = os.path.join(cz_dir, f"Carta_Beneficiarios_{eas.replace(' ', '_')}.docx")
    doc.save(output_path)
    print(f"Carta generada exitosamente: {output_path}")


print(f"Proceso finalizado. Cartas generadas para todas las EAS en {CZ} con enlace {ENLACE}.")